import base64
import io
import ipaddress
import logging
import socket
from urllib.parse import urljoin
from urllib.parse import urlparse

import requests
from bs4 import BeautifulSoup
from docx import Document
from pypdf import PdfReader


logger = logging.getLogger("helios-ingestion")

MAX_REMOTE_BYTES = 2_000_000
MAX_REDIRECTS = 3
MAX_GITHUB_TREE_ENTRIES = 300
USER_AGENT = "HELIOS-Source-Ingestion/1.0"


def _validate_public_url(
    url: str,
    allowed_hosts: set[str] | None = None
) -> str:
    parsed = urlparse(
        str(url or "").strip()
    )
    if parsed.scheme not in {
        "http",
        "https"
    }:
        raise ValueError(
            "Only HTTP and HTTPS URLs are supported."
        )
    hostname = (
        parsed.hostname or ""
    ).lower()
    if not hostname:
        raise ValueError(
            "URL hostname is required."
        )
    if allowed_hosts and hostname not in allowed_hosts:
        raise ValueError(
            "URL host is not allowed for this source type."
        )
    if hostname == "localhost" or hostname.endswith(
        ".localhost"
    ):
        raise ValueError(
            "Local URLs are not allowed."
        )
    try:
        addresses = {
            item[4][0]
            for item in socket.getaddrinfo(
                hostname,
                parsed.port or (
                    443
                    if parsed.scheme == "https"
                    else 80
                ),
                type=socket.SOCK_STREAM
            )
        }
    except socket.gaierror as error:
        raise ValueError(
            f"Could not resolve URL hostname: {hostname}"
        ) from error
    if not addresses:
        raise ValueError(
            "URL hostname did not resolve."
        )
    for address in addresses:
        ip = ipaddress.ip_address(
            address
        )
        if not ip.is_global:
            raise ValueError(
                "Private, loopback, and link-local URLs are not allowed."
            )
    return parsed.geturl()


def _read_limited_response(
    response: requests.Response,
    max_bytes: int = MAX_REMOTE_BYTES
) -> bytes:
    content_length = response.headers.get(
        "content-length"
    )
    if content_length:
        try:
            if int(
                content_length
            ) > max_bytes:
                raise ValueError(
                    "Remote source exceeds the ingestion size limit."
                )
        except ValueError as error:
            if "exceeds" in str(
                error
            ):
                raise

    content = bytearray()
    for chunk in response.iter_content(
        chunk_size=65536
    ):
        if not chunk:
            continue
        content.extend(
            chunk
        )
        if len(
            content
        ) > max_bytes:
            raise ValueError(
                "Remote source exceeds the ingestion size limit."
            )
    return bytes(
        content
    )


def _get_public_url(
    url: str,
    *,
    allowed_hosts: set[str] | None = None,
    max_bytes: int = MAX_REMOTE_BYTES
) -> requests.Response:
    current_url = str(
        url
    )
    for _ in range(
        MAX_REDIRECTS + 1
    ):
        current_url = _validate_public_url(
            current_url,
            allowed_hosts
        )
        response = requests.get(
            current_url,
            timeout=10,
            allow_redirects=False,
            stream=True,
            headers={
                "User-Agent": USER_AGENT,
                "Accept": "text/html,application/json,text/plain;q=0.9,*/*;q=0.1"
            }
        )
        if response.is_redirect or response.is_permanent_redirect:
            location = response.headers.get(
                "location"
            )
            if not location:
                raise ValueError(
                    "Remote source returned an invalid redirect."
                )
            current_url = urljoin(
                current_url,
                location
            )
            continue
        response.raise_for_status()
        response._content = _read_limited_response(
            response,
            max_bytes
        )
        response.encoding = response.encoding or "utf-8"
        return response
    raise ValueError(
        "Remote source exceeded the redirect limit."
    )


def extract_pdf_text(
    file_bytes: bytes
) -> str:
    """Extract text from PDF bytes."""
    try:
        reader = PdfReader(
            io.BytesIO(
                file_bytes
            )
        )
        text = "\n".join(
            page.extract_text() or ""
            for page in reader.pages
        )
        return text.strip()
    except Exception as error:
        logger.exception(
            "Failed to extract PDF text."
        )
        raise ValueError(
            "PDF extraction failed."
        ) from error


def extract_docx_text(
    file_bytes: bytes
) -> str:
    """Extract paragraphs and table cells from DOCX bytes."""
    try:
        document = Document(
            io.BytesIO(
                file_bytes
            )
        )
        blocks = [
            paragraph.text.strip()
            for paragraph in document.paragraphs
            if paragraph.text.strip()
        ]
        for table in document.tables:
            for row in table.rows:
                cells = [
                    cell.text.strip()
                    for cell in row.cells
                    if cell.text.strip()
                ]
                if cells:
                    blocks.append(
                        " | ".join(
                            cells
                        )
                    )
        return "\n".join(
            blocks
        ).strip()
    except Exception as error:
        logger.exception(
            "Failed to extract DOCX text."
        )
        raise ValueError(
            "DOCX extraction failed."
        ) from error


def extract_web_text(
    url: str
) -> str:
    """Extract visible text from a public web URL."""
    response = _get_public_url(
        url
    )
    soup = BeautifulSoup(
        response.text,
        "html.parser"
    )
    for element in soup(
        [
            "script",
            "style",
            "noscript"
        ]
    ):
        element.decompose()
    lines = (
        line.strip()
        for line in soup.get_text(
            separator="\n"
        ).splitlines()
    )
    return "\n".join(
        line
        for line in lines
        if line
    ).strip()


def _github_repo_parts(
    repo_url: str
) -> tuple[str, str]:
    parsed = urlparse(
        _validate_public_url(
            repo_url,
            {
                "github.com"
            }
        )
    )
    parts = [
        part
        for part in parsed.path.split(
            "/"
        )
        if part
    ]
    if len(
        parts
    ) != 2:
        raise ValueError(
            "GitHub source must be a repository URL."
        )
    owner, repo = parts
    if repo.endswith(
        ".git"
    ):
        repo = repo[:-4]
    if not owner or not repo:
        raise ValueError(
            "GitHub source must be a repository URL."
        )
    return owner, repo


def extract_github_repo(
    repo_url: str
) -> str:
    """Extract a public GitHub repository README and file structure."""
    owner, repo = _github_repo_parts(
        repo_url
    )
    api_hosts = {
        "api.github.com"
    }
    repo_response = _get_public_url(
        f"https://api.github.com/repos/{owner}/{repo}",
        allowed_hosts=api_hosts
    )
    repo_data = repo_response.json()
    default_branch = repo_data.get(
        "default_branch",
        "main"
    )

    tree_response = _get_public_url(
        f"https://api.github.com/repos/{owner}/{repo}/git/trees/{default_branch}?recursive=1",
        allowed_hosts=api_hosts
    )
    tree_data = tree_response.json()
    paths = [
        item.get(
            "path"
        )
        for item in tree_data.get(
            "tree",
            []
        )
        if item.get(
            "type"
        )
        == "blob"
        and item.get(
            "path"
        )
    ][
        :MAX_GITHUB_TREE_ENTRIES
    ]

    readme_text = ""
    try:
        readme_response = _get_public_url(
            f"https://api.github.com/repos/{owner}/{repo}/readme",
            allowed_hosts=api_hosts
        )
        readme_data = readme_response.json()
        if readme_data.get(
            "encoding"
        ) == "base64":
            readme_text = base64.b64decode(
                readme_data.get(
                    "content",
                    ""
                )
            ).decode(
                "utf-8",
                errors="replace"
            )
    except Exception:
        logger.info(
            "README not available for GitHub repository %s/%s.",
            owner,
            repo
        )

    structure = "\n".join(
        f"- {path}"
        for path in paths
    )
    return (
        f"GitHub Repository: https://github.com/{owner}/{repo}\n"
        f"Default Branch: {default_branch}\n\n"
        f"Repository Structure:\n{structure or '- No files returned'}\n\n"
        f"README:\n{readme_text or 'README not available.'}"
    ).strip()
