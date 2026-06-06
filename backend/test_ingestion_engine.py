import socket
from io import BytesIO

import pytest
from docx import Document

from core import ingestion_engine


def test_web_ingestion_blocks_private_addresses(
    monkeypatch
):
    monkeypatch.setattr(
        socket,
        "getaddrinfo",
        lambda *args, **kwargs: [
            (
                socket.AF_INET,
                socket.SOCK_STREAM,
                6,
                "",
                (
                    "127.0.0.1",
                    80
                )
            )
        ]
    )

    with pytest.raises(
        ValueError,
        match="Private"
    ):
        ingestion_engine.extract_web_text(
            "http://example.test/private"
        )


def test_github_repo_requires_repository_url(
    monkeypatch
):
    monkeypatch.setattr(
        socket,
        "getaddrinfo",
        lambda *args, **kwargs: [
            (
                socket.AF_INET,
                socket.SOCK_STREAM,
                6,
                "",
                (
                    "140.82.112.4",
                    443
                )
            )
        ]
    )

    with pytest.raises(
        ValueError,
        match="repository URL"
    ):
        ingestion_engine.extract_github_repo(
            "https://github.com/openai"
        )


def test_docx_ingestion_extracts_paragraphs_and_tables():
    document = Document()
    document.add_paragraph(
        "HELIOS document intelligence"
    )
    table = document.add_table(
        rows=1,
        cols=1
    )
    table.cell(
        0,
        0
    ).text = "Grounded source row"
    output = BytesIO()
    document.save(
        output
    )

    text = ingestion_engine.extract_docx_text(
        output.getvalue()
    )

    assert "HELIOS document intelligence" in text
    assert "Grounded source row" in text
